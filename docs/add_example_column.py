"""为 dflash_command_reference.xlsx 的"子命令总览" sheet 新增"示例"列。"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

XLSX = r"d:\workspace\tas-dflash_-erase\docs\dflash_command_reference.xlsx"

EXAMPLES = {
    "erase":   "dflash erase --addr AF000000 --sectors 2 --backup erase_bak.hex",
    "write":   "dflash write --file firmware.hex --verify",
    "rewrite": "dflash rewrite --addr AF004004 --data 1122334455667788 --verify --backup",
    "read":    "dflash read --addr AF000000 --length 2000 --file output.hex",
    "restore": "dflash restore --file backup.hex --verify",
    "list":    "dflash list",
    "ucb":     "dflash ucb read --addr AF400000 --length 200",
}

wb = load_workbook(XLSX)
ws = wb["子命令总览"]

new_col = ws.max_column + 1
header_src = ws.cell(row=1, column=1)
last_data_src = ws.cell(row=2, column=ws.max_column)  # 复用现有数据单元格的样式

# 写入标题
header_cell = ws.cell(row=1, column=new_col, value="示例")
header_cell.font = copy(header_src.font)
header_cell.fill = copy(header_src.fill)
header_cell.alignment = copy(header_src.alignment)
header_cell.border = copy(header_src.border)

# 写入数据行
for row in range(2, ws.max_row + 1):
    sub_cmd = ws.cell(row=row, column=1).value
    example = EXAMPLES.get(sub_cmd, "")
    cell = ws.cell(row=row, column=new_col, value=example)
    cell.font = copy(last_data_src.font)
    cell.alignment = copy(last_data_src.alignment)
    cell.border = copy(last_data_src.border)
    if last_data_src.fill and last_data_src.fill.patternType:
        cell.fill = copy(last_data_src.fill)

# 设置列宽
from openpyxl.utils import get_column_letter
ws.column_dimensions[get_column_letter(new_col)].width = 70

wb.save(XLSX)
print(f"已添加列 {get_column_letter(new_col)} '示例'，保存至 {XLSX}")
